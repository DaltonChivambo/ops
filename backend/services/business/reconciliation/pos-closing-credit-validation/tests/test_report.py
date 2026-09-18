"""O relatório Excel: o que cai em cada folha, e se as folhas fecham entre si.

O DOP lê o relatório, não a API — um número errado aqui é um número errado no
fecho do departamento, mesmo com a reconciliação certa. Por isso os testes não
olham só para células soltas: obrigam as três folhas a baterem umas com as
outras, que é a conferência que o operador faria à mão.

Um cenário só, com um fecho de cada situação, a passar pela reconciliação real.
"""

from dataclasses import asdict
from datetime import date
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from typing import Any

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.domain.models import BankaCredit, BankaMovement, PosInfo, SimoClosing
from app.domain.reconciliation import reconcile
from app.domain.vocabulary import ClosingType, Validation
from app.infrastructure.excel.report import build_workbook

DAY = date(2026, 6, 23)

# Colunas (a A fica vazia em todas as folhas).
DETAILS_SIMO, DETAILS_BANKA, DETAILS_VALIDATION = 9, 12, 14
PENDING_SIMO, PENDING_BANKA, PENDING_VALIDATION = 10, 13, 14

MATCH = "Crédito Confere"
MISMATCH = "Fecho creditado incorrectamente"
MISSING = "Fecho Não Creditado_aguarda tratamento da SIMO"
DUPLICATED = "Períodos repetidos_analisar individualmente"


def _closing(pos_id: str, period: int, total: str, ops: int = 1) -> SimoClosing:
    return SimoClosing(pos_id, period, DAY, ops, Decimal(total))


def _credit(*amounts: str) -> BankaCredit:
    movements = [BankaMovement(DAY, Decimal(m), "P24-Fecho TPA") for m in amounts]
    total = sum((m.amount for m in movements), Decimal(0))
    return BankaCredit(total, DAY, "P24-Fecho TPA", movements)


def _workbook(reconciled_pos: str | None = None) -> Any:
    """O livro do cenário. `reconciled_pos` concilia o 1º fecho desse POS, como o ecrã faria."""
    return load_workbook(BytesIO(_workbook_bytes(reconciled_pos)))


def _workbook_bytes(reconciled_pos: str | None = None) -> bytes:
    pos_ids = ["200001", "200002", "200003", "200004", "200005", "200006", "200007"]
    pos_list = {p: PosInfo(f"Comerciante {p}", f"ACC{p}", ClosingType.D) for p in pos_ids}
    closings = [
        _closing("200001", 101, "100.00"),  # confere
        _closing("200002", 102, "200.00"),  # incorrecto: Banka 190
        _closing("200003", 103, "300.00"),  # não creditado
        _closing("200004", 104, "50.00", ops=1),  # duplicado, por tratar
        _closing("200004", 104, "70.00", ops=2),
        _closing("200005", 105, "40.00", ops=1),  # duplicado, já regularizado
        _closing("200005", 105, "60.00", ops=2),
        _closing("200006", 106, "80.00"),  # um fecho, dois movimentos que não batem
        _closing("200007", 107, "0.00"),  # zerado
    ]
    credits = {
        "200001101": _credit("100.00"),
        "200002102": _credit("190.00"),
        "200004104": _credit("50.00", "70.00"),
        "200005105": _credit("40.00", "60.00"),
        "200006106": _credit("30.00", "60.00"),
    }
    result = reconcile(pos_list, closings, credits)
    if reconciled_pos is not None:
        detail = next(d for d in result.details if d.pos_id == reconciled_pos)
        detail.validation = Validation.MATCH
        detail.difference = Decimal(0)

    execution = SimpleNamespace(
        summary=result.summary.to_json_dict(),
        period_start=result.period_start,
        period_end=result.period_end,
    )
    cases = [
        SimpleNamespace(
            **asdict(case),
            status="resolved" if case.pos_id == "200005" else "pending",
            e_ticket=None,
            resolved_at=None,
        )
        for case in result.cases
    ]
    return build_workbook(execution, result.details, cases)


@pytest.fixture(scope="module")
def workbook() -> Any:
    return _workbook()


def _row(sheet: Worksheet, row: int, columns: int) -> list[Any]:
    return [sheet.cell(row=row, column=2 + offset).value for offset in range(columns)]


def _data_rows(sheet: Worksheet) -> list[tuple[Any, ...]]:
    return [
        row
        for row in sheet.iter_rows(min_row=3, values_only=True)
        if any(value is not None for value in row)
    ]


def _column_sum(rows: list[tuple[Any, ...]], column: int) -> float:
    """Soma só os números — «n.a» é a célula que de propósito não entra na soma."""
    return sum(row[column - 1] for row in rows if isinstance(row[column - 1], int | float))


# ─── Cada folha ──────────────────────────────────────────────────────────────


def test_summary_has_one_row_per_cause_like_execution(workbook: Any) -> None:
    """Um período duplicado não é um crédito incorrecto — não entra nessa linha."""
    summary = workbook["Resumo"]

    assert _row(summary, 7, 5) == pytest.approx([MISMATCH, 1, 200, 190, -10])
    assert _row(summary, 8, 5) == pytest.approx([MISSING, 1, 300, 0, -300])
    # Duplicados: 2 + 2 fechos em chaves com vários fechos, 1 com vários movimentos.
    # O Banka creditou 310 nestas chaves, mas a SIMO só fechou 300: o excesso não
    # tem fecho que o pague e fica de fora — ver `compute_summary`.
    assert _row(summary, 9, 5) == pytest.approx([DUPLICATED, 5, 300, 300, 0])
    assert _row(summary, 10, 5) == pytest.approx([MATCH, 1, 100, 100, 0])
    assert _row(summary, 11, 5) == pytest.approx(["Total", 8, 900, 590, -310])


def test_summary_explains_file_lines_and_repeated_rows(workbook: Any) -> None:
    """Quantos fechos veio no ficheiro, e quantos desses são duplicados na SIMO."""
    summary = workbook["Resumo"]

    assert summary["B27"].value == "Linhas dos ficheiros"
    labels = {
        summary.cell(row=row, column=2).value: summary.cell(row=row, column=3).value
        for row in range(30, 33)
    }
    assert labels["Fechos no ficheiro da SIMO"] == summary["C11"].value + 1  # + o zerado
    assert labels["Dos quais linhas repetidas"] == 0


def test_pending_summary_counts_closings_without_double_counting_duplicates(workbook: Any) -> None:
    """Um caso de período duplicado junta vários fechos — conta cada fecho uma vez."""
    summary = workbook["Resumo"]

    assert summary["B15"].value == "Total Casos Pendentes na SIMO"
    assert _row(summary, 19, 3) == pytest.approx(["Fecho Regularizado", 2, 100])
    assert _row(summary, 20, 3) == pytest.approx([MISMATCH, 1, 200])
    assert _row(summary, 21, 3) == pytest.approx([MISSING, 1, 300])
    # O 200005 já foi regularizado: ficam os dois fechos do 200004 e o do 200006.
    assert _row(summary, 22, 3) == pytest.approx([DUPLICATED, 3, 200])
    assert _row(summary, 23, 3) == pytest.approx(["Total", 7, 800])


def test_reconciled_closing_of_open_case_counts_as_resolved() -> None:
    """A chave 200004 só conciliada em parte: o caso continua aberto, mas o fecho
    conciliado já não está por tratar — e o relatório gera-se, em vez de rebentar."""
    summary = _workbook(reconciled_pos="200004")["Resumo"]

    assert _row(summary, 19, 3) == pytest.approx(["Fecho Regularizado", 3, 150])
    assert _row(summary, 22, 3) == pytest.approx([DUPLICATED, 2, 150])
    assert _row(summary, 23, 3) == pytest.approx(["Total", 7, 800])


def test_pending_sheet_only_has_open_items_with_right_cause(workbook: Any) -> None:
    rows = _data_rows(workbook["Total Casos Pendentes na SIMO"])

    assert [(row[1], row[PENDING_VALIDATION - 1]) for row in rows] == [
        ("200002", MISMATCH),
        ("200003", MISSING),
        ("200004", DUPLICATED),
        ("200004", DUPLICATED),
        ("200006", DUPLICATED),
    ]


def test_credit_of_multi_closing_key_only_on_first_row(workbook: Any) -> None:
    """Repetido em cada fecho, o crédito da chave somava a dobrar — o erro do VLOOKUP."""
    details = [r for r in _data_rows(workbook["Detalhes Validacao"]) if r[1] == "200004"]
    pending = [r for r in _data_rows(workbook["Total Casos Pendentes na SIMO"]) if r[1] == "200004"]

    assert [r[DETAILS_BANKA - 1] for r in details] == [120, "n.a"]
    assert [r[PENDING_BANKA - 1] for r in pending] == [120, "n.a"]


# ─── As folhas fecham entre si ───────────────────────────────────────────────


def test_details_sum_matches_summary_total(workbook: Any) -> None:
    """A SIMO fecha entre as folhas. O Banka não, e é de propósito.

    Os detalhes mostram o crédito da chave tal como o banco o fez — é o que o
    operador precisa de ver. O resumo só conta esse crédito até onde há fecho
    que ele possa pagar, por isso é menor ou igual, nunca maior.
    """
    summary = workbook["Resumo"]
    details = _data_rows(workbook["Detalhes Validacao"])

    assert _column_sum(details, DETAILS_SIMO) == pytest.approx(summary["D11"].value)
    assert _column_sum(details, DETAILS_BANKA) >= summary["E11"].value


def test_details_count_per_cause_matches_summary(workbook: Any) -> None:
    summary = workbook["Resumo"]
    details = _data_rows(workbook["Detalhes Validacao"])

    for row in range(7, 11):
        label = summary.cell(row=row, column=2).value
        assert (
            sum(1 for r in details if r[DETAILS_VALIDATION - 1] == label)
            == summary.cell(row=row, column=3).value
        )


def test_pending_sheet_matches_summary_open_rows(workbook: Any) -> None:
    summary = workbook["Resumo"]
    pending = _data_rows(workbook["Total Casos Pendentes na SIMO"])

    for row in range(20, 23):
        label = summary.cell(row=row, column=2).value
        of_cause = [r for r in pending if r[PENDING_VALIDATION - 1] == label]
        assert len(of_cause) == summary.cell(row=row, column=3).value
        assert _column_sum(of_cause, PENDING_SIMO) == pytest.approx(
            summary.cell(row=row, column=4).value
        )


def test_resolved_plus_pending_equals_all_non_matching(workbook: Any) -> None:
    summary = workbook["Resumo"]

    non_matching = sum(summary.cell(row=row, column=3).value for row in range(7, 10))
    amount = sum(summary.cell(row=row, column=4).value for row in range(7, 10))
    assert summary["C23"].value == non_matching
    assert summary["D23"].value == pytest.approx(amount)


# ─── Aspecto ─────────────────────────────────────────────────────────────────


def test_summary_carries_logo_and_context_line(workbook: Any) -> None:
    summary = workbook["Resumo"]

    assert len(summary._images) == 1
    assert summary["B4"].value.startswith("Departamento de Meios de Pagamento e Canais")


def test_data_sheets_have_filters_frozen_header_and_state_colours(workbook: Any) -> None:
    for name in ("Detalhes Validacao", "Total Casos Pendentes na SIMO"):
        sheet = workbook[name]

        assert sheet.auto_filter.ref.startswith("B2:")
        assert sheet.freeze_panes == "C3"
        assert sheet.print_title_rows == "$2:$2"
        # Linhas alternadas e cor da validação por formatação condicional, não célula a célula.
        assert len(sheet.conditional_formatting) >= 2


def test_details_follow_description_period_within_each_pos() -> None:
    """Dentro do POS, os fechos vêm pelo período do fim do descritivo, crescente."""
    from types import SimpleNamespace as Row

    def closing(pos_id: str, description: str | None, period: int) -> Row:
        return Row(
            pos_id=pos_id,
            closing_description=description,
            period=period,
            simo_closing_date=DAY,
            operation_number=1,
        )

    from app.infrastructure.excel.report import _details_order

    rows = [
        closing("263073", "P24-Fecho TPA 0000263073 - 003", 3),
        closing("263073", "P24-Fecho TPA 0000263073 - 001", 1001),
        closing("100001", None, 2002),
        closing("263073", "P24-Fecho TPA 0000263073 - 002", 2),
    ]

    ordered = sorted(rows, key=_details_order)

    assert [(r.pos_id, r.closing_description) for r in ordered] == [
        ("100001", None),
        ("263073", "P24-Fecho TPA 0000263073 - 001"),
        ("263073", "P24-Fecho TPA 0000263073 - 002"),
        ("263073", "P24-Fecho TPA 0000263073 - 003"),
    ]


def test_details_mark_case_columns_not_applicable_when_credit_matches(workbook: Any) -> None:
    details = workbook["Detalhes Validacao"]

    for row in _data_rows(details):
        if row[DETAILS_VALIDATION - 1] == MATCH:
            assert row[16 - 1] == "n.a"  # e-Ticket (coluna P)
            assert row[17 - 1] == "n.a"  # Data Reg. (coluna Q)


def test_details_identity_headers_have_their_own_colour(workbook: Any) -> None:
    details = workbook["Detalhes Validacao"]
    fills = {
        details.cell(row=2, column=column).value: details.cell(
            row=2, column=column
        ).fill.fgColor.rgb
        for column in range(2, 18)
    }

    assert fills["POS ID"] == fills["Data Reg."] == "FF1F2A44"
    assert fills["Total Fecho \nSIMO"] == fills["Validação"] == "FFC00000"


def test_digit_columns_do_not_warn_number_stored_as_text() -> None:
    """POS ID, conta e chave são texto de propósito: o Excel não os marca com o aviso."""
    from zipfile import ZipFile

    archive = ZipFile(BytesIO(_workbook_bytes()))
    sheets = [n for n in archive.namelist() if n.startswith("xl/worksheets/sheet")]
    marked = [n for n in sheets if b'numberStoredAsText="1"' in archive.read(n)]

    # As duas folhas de dados; o Resumo não tem colunas destas.
    assert len(marked) == 2


def test_pending_identity_headers_have_their_own_colour(workbook: Any) -> None:
    pending = workbook["Total Casos Pendentes na SIMO"]
    headers = [pending.cell(row=2, column=column) for column in range(2, 16)]
    fills = {cell.value: cell.fill.fgColor.rgb for cell in headers}

    assert fills["POS ID"] == fills["Data Reg."] == "FF1F2A44"
    assert fills["Total Fecho \nSIMO"] == fills["Validação"] == "FFC00000"


def test_simo_duplicate_counts_with_original_validation_and_is_marked() -> None:
    """A cópia do export vem a seguir à original, com a mesma validação, marcada."""
    pos_list = {"300001": PosInfo("Comerciante", "ACC", ClosingType.D)}
    closings = [_closing("300001", 101, "100.00"), _closing("300001", 101, "100.00")]
    result = reconcile(pos_list, closings, {"300001101": _credit("100.00")})
    execution = SimpleNamespace(
        summary=result.summary.to_json_dict(),
        period_start=result.period_start,
        period_end=result.period_end,
    )
    book = load_workbook(BytesIO(build_workbook(execution, result.details, [])))
    details = book["Detalhes Validacao"]

    rows = _data_rows(details)
    assert [row[DETAILS_VALIDATION - 1] for row in rows] == [MATCH, MATCH]
    assert rows[1][DETAILS_BANKA - 1] == "n.a"
    summary = book["Resumo"]
    assert summary["C10"].value == 1  # confere conta o fecho uma vez
    assert summary["D10"].value == 100
    assert summary["B11"].value == "Linhas repetidas na SIMO"
    assert summary["C11"].value == 1
    assert summary["B12"].value == "Total"
    assert summary["C12"].value == 2
    assert summary["D12"].value == 100  # o montante conta uma vez
    marked = [str(cf.sqref) for cf in details.conditional_formatting]
    assert "E3:F4" in marked
