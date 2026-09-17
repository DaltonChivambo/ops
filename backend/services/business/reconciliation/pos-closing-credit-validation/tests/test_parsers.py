"""Leitura dos ficheiros de entrada: o que varia entre exports e não pode partir.

Os exports do Portal SIMO não vêm sempre com a mesma margem à esquerda — os do
departamento começam ora na coluna B ora na C, com as linhas de cima vazias. Por
isso as colunas resolvem-se pelo NOME do cabeçalho e não pela posição: com
posições fixas, um export deslocado uma casa lia «Id Comerciante» como POS Id
(nenhum casava com a Lista de POS) e «Período POS» como data — e um período como
4920, lido como serial do Excel, dava uma data de 1913.

Do lado do Banka, o que pode partir é o próprio extracto: juntado a partir de
dois que se sobrepõem, traz o dia da junção duas vezes, e cada movimento desse
dia somava a dobrar.

Sem `.xlsx` de lado nenhum: as folhas são construídas aqui, célula a célula.
"""

from datetime import date
from decimal import Decimal
from io import BytesIO
from typing import Any

import pytest
from openpyxl import Workbook

from app.infrastructure.excel.parsers import parse_banka_credits, parse_simo_closings
from app.infrastructure.excel.workbook import cell_date

HEADERS = [
    "Id Comerciante",
    "POS Id",
    "Período POS",
    "Data Fecho",
    "Nº Operaç.",
    "Total Fecho",
    "Downl.",
]
ROW = [132295, 237958, 4920, date(2026, 8, 20), 1, "12.999,99"]


def _simo_sheet(*, first_column: int, first_row: int) -> BytesIO:
    """Um export da SIMO com o cabeçalho onde o chamador o puser."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    for offset, title in enumerate(HEADERS):
        sheet.cell(row=first_row, column=first_column + offset, value=title)
    for offset, value in enumerate(ROW):
        sheet.cell(row=first_row + 1, column=first_column + offset, value=value)

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


@pytest.mark.parametrize(
    ("first_column", "first_row"),
    [(2, 1), (3, 3), (1, 1)],
    ids=["comeca-em-B", "comeca-em-C-na-linha-3", "comeca-em-A"],
)
def test_closing_is_the_same_wherever_the_export_starts(first_column: int, first_row: int) -> None:
    (closing,) = parse_simo_closings(
        _simo_sheet(first_column=first_column, first_row=first_row),
        "simo.xlsx",
    )

    assert closing.pos_id == "237958"  # e não 132295, o Id do Comerciante
    assert closing.period == 4920
    assert closing.closing_date == date(2026, 8, 20)  # e não uma data de 1913
    assert closing.operation_number == 1
    assert closing.total == Decimal("12999.99")


@pytest.mark.parametrize("number", [1, 70, 4920, 0])
def test_number_that_is_not_a_date_does_not_become_a_date(number: int) -> None:
    """Um período lido por engano na coluna da data não pode virar 1899.

    Era o que punha uma execução de Agosto de 2026 a dizer-se «31 de Dezembro a
    10 de Março»: o serial 1 do Excel é 31/12/1899 e o 70 é 10/03/1900.
    """
    assert cell_date(number) is None


def test_real_excel_serial_is_still_a_date() -> None:
    assert cell_date(46254) == date(2026, 8, 20)


# ─── Créditos do Banka: cada movimento uma vez ───────────────────────────────

BANKA_HEADERS = ["DATA_SISTEMA", "N_DOCUMENTO", "DESCRITIVO_MOV", "VALOR_TRANSACAO"]
DESCRIPTION = "P24-Fecho TPA 0000221337 - 486"
KEY = "221337486"


def _banka_sheet(rows: list[list[Any]], headers: list[str] = BANKA_HEADERS) -> BytesIO:
    """Um export do MIS: título, linha vazia e o cabeçalho na terceira, como o real."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "FECHO_POS"
    sheet.append(["FECHO_POS"])
    sheet.append([])
    sheet.append(headers)
    for row in rows:
        sheet.append(row)

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


def test_repeated_movement_in_extract_counts_once() -> None:
    """O mesmo N_DOCUMENTO duas vezes é o mesmo movimento — o dia juntado a dobrar."""
    row = ["19/08/2026", 761076120, DESCRIPTION, 2278]

    credits, discarded = parse_banka_credits(_banka_sheet([row, list(row)]), "banka.xlsx")

    assert discarded == 1
    assert len(credits[KEY].movements) == 1
    assert credits[KEY].amount == Decimal("2278")


def test_equal_movements_with_different_documents_are_two_credits() -> None:
    """Mesmo dia, mesmo valor, mesma chave — mas dois documentos: não se apaga nenhum."""
    rows = [
        ["19/08/2026", 761076120, DESCRIPTION, 2278],
        ["19/08/2026", 761076121, DESCRIPTION, 2278],
    ]

    credits, discarded = parse_banka_credits(_banka_sheet(rows), "banka.xlsx")

    assert discarded == 0
    assert len(credits[KEY].movements) == 2
    assert credits[KEY].amount == Decimal("4556")


def test_without_document_number_only_identical_rows_collapse() -> None:
    headers = ["DATA_SISTEMA", "DESCRITIVO_MOV", "VALOR_TRANSACAO"]
    rows = [
        ["19/08/2026", DESCRIPTION, 2278],
        ["19/08/2026", DESCRIPTION, 2278],  # cópia exacta
        ["20/08/2026", DESCRIPTION, 2278],  # outro dia: outro movimento
    ]

    credits, discarded = parse_banka_credits(_banka_sheet(rows, headers), "banka.xlsx")

    assert discarded == 1
    assert [m.date for m in credits[KEY].movements] == [date(2026, 8, 19), date(2026, 8, 20)]
