"""O aspecto do relatório: cores, cabeçalhos, linhas, logótipo e impressão.

À parte do `report.py`, que decide o que vai em cada célula. Aqui decide-se só
como se vê — e sempre de forma que as ~30 mil linhas de uma execução real não
pesem: as folhas de detalhe não pintam célula a célula, usam formatação
condicional (linhas alternadas, cor da validação, diferença a vermelho), que o
Excel aplica ao abrir e o ficheiro guarda uma vez por intervalo.
"""

import re
import zipfile
from collections.abc import Sequence
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.drawing.image import Image
from openpyxl.formatting.rule import Rule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

LOGO = Path(__file__).parent / "assets" / "mozaops-logo.png"
LOGO_HEIGHT_PX = 46

# ─── Paleta ──────────────────────────────────────────────────────────────────

MOZA_RED = "FFC00000"
# Cabeçalho das colunas de identificação e datas — quem, quando — para se
# distinguirem das de montantes e validação, que ficam a vermelho.
MOZA_NAVY = "FF1F2A44"
INK = "FF1F2937"
MUTED = "FF6B7280"
LINE = "FFE5E7EB"
ZEBRA = "FFF9FAFB"
TOTAL = "FFF3F4F6"

# Tom de cada estado: fundo claro e texto escuro da mesma cor — os mesmos do ecrã.
STATE_TONES = {
    "match": ("FFECFDF5", "FF047857"),
    "mismatch": ("FFFEF2F2", "FFB91C1C"),
    "missing": ("FFEEF2FF", "FF3730A3"),
    "duplicated": ("FFFFFBEB", "FFB45309"),
    "zero": ("FFF3F4F6", "FF6B7280"),
    "resolved": ("FFECFDF5", "FF047857"),
    # Fecho repetido no export: o vermelho claro dos duplicados.
    "repeated": ("FFFEE2E2", "FFB91C1C"),
}

_THIN_LINE = Side(style="thin", color=LINE)
_RED_LINE = Side(style="medium", color=MOZA_RED)

HEADER_FILL = PatternFill("solid", fgColor=MOZA_RED)
IDENTITY_HEADER_FILL = PatternFill("solid", fgColor=MOZA_NAVY)
CENTER = Alignment(horizontal="center", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")
HEADER_FONT = Font(bold=True, color="FFFFFFFF", size=10)
HEADER_ALIGN = Alignment(vertical="center", horizontal="center", wrap_text=True)
HEADER_BORDER = Border(
    left=Side(style="thin", color="FFFFFFFF"), right=Side(style="thin", color="FFFFFFFF")
)

ROW_BORDER = Border(bottom=_THIN_LINE)
TOTAL_FILL = PatternFill("solid", fgColor=TOTAL)
TOTAL_FONT = Font(bold=True, color=INK)
TOTAL_BORDER = Border(top=_RED_LINE, bottom=_THIN_LINE)


def header_cell(cell: Cell | MergedCell, *, identity: bool = False) -> None:
    cell.fill = IDENTITY_HEADER_FILL if identity else HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = HEADER_ALIGN
    cell.border = HEADER_BORDER


def state_cell(cell: Cell | MergedCell, state: str) -> None:
    """A célula do rótulo de um estado, no Resumo: fundo e texto do tom do estado."""
    fill, color = STATE_TONES[state]
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(bold=True, color=color)


# ─── Folhas ──────────────────────────────────────────────────────────────────


def document(workbook: Workbook, title: str) -> None:
    workbook.properties.title = title
    workbook.properties.creator = "MozaOps"
    workbook.properties.company = "Moza Banco · DOP"


def cover(sheet: Worksheet, title: str, subtitle: str) -> None:
    """O topo do Resumo: logótipo, título e a linha de contexto, sem grelha."""
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 40
    sheet.row_dimensions[2].height = 8
    if LOGO.exists():
        logo = Image(str(LOGO))
        ratio = LOGO_HEIGHT_PX / logo.height
        logo.height = LOGO_HEIGHT_PX
        logo.width = round(logo.width * ratio)
        sheet.add_image(logo, "B1")

    sheet["B3"] = title
    sheet["B3"].font = Font(bold=True, size=16, color=INK)
    sheet.row_dimensions[3].height = 24
    sheet["B4"] = subtitle
    sheet["B4"].font = Font(size=10, color=MUTED)


def section_title(sheet: Worksheet, cell: str, last_column: str) -> None:
    """Título de um bloco do Resumo, sublinhado a vermelho até à última coluna."""
    row = int("".join(ch for ch in cell if ch.isdigit()))
    sheet[cell].font = Font(bold=True, size=13, color=INK)
    sheet.row_dimensions[row].height = 22
    for column in range(2, column_index_from_string(last_column) + 1):
        sheet.cell(row=row, column=column).border = Border(bottom=_RED_LINE)


def summary_row(sheet: Worksheet, row: int, columns: int, *, total: bool = False) -> None:
    sheet.row_dimensions[row].height = 20
    for column in range(2, 2 + columns):
        cell = sheet.cell(row=row, column=column)
        cell.alignment = Alignment(vertical="center", indent=1 if column == 2 else 0)
        if total:
            cell.fill = TOTAL_FILL
            cell.font = TOTAL_FONT
            cell.border = TOTAL_BORDER
        else:
            cell.border = ROW_BORDER


def empty_message(cell: Cell | MergedCell) -> None:
    """A linha única de uma folha sem dados."""
    cell.font = Font(italic=True, color=MUTED)
    cell.alignment = Alignment(horizontal="center", vertical="center")


def header_row(sheet: Worksheet, row: int) -> None:
    sheet.row_dimensions[row].height = 34


def data_sheet(
    sheet: Worksheet,
    *,
    header: int,
    last_row: int,
    columns: int,
    validation_column: str,
    difference_column: str | None,
    tab_color: str,
    marked: Sequence[str] = (),
) -> None:
    """Folha de detalhe: filtros, colunas fixas, linhas alternadas e cores, pronta a imprimir."""
    sheet.sheet_properties.tabColor = tab_color
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 10
    header_row(sheet, header)

    last_column = get_column_letter(1 + columns)
    first_data = header + 1
    if last_row >= first_data:
        data = f"B{first_data}:{last_column}{last_row}"
        sheet.auto_filter.ref = f"B{header}:{last_column}{last_row}"
        if marked:
            # Antes das linhas alternadas: a primeira regra ganha no fundo, e o
            # vermelho não pode ficar tapado pelo cinzento das linhas pares.
            red_fill, red_text = STATE_TONES["repeated"]
            sheet.conditional_formatting.add(
                " ".join(marked),
                _when(
                    formula=["TRUE"],
                    fill=PatternFill(bgColor=red_fill, fill_type="solid"),
                    font=Font(bold=True, color=red_text),
                    stopIfTrue=True,
                ),
            )
        # Linhas alternadas e a risca fina entre linhas, numa regra para o intervalo inteiro.
        sheet.conditional_formatting.add(
            data,
            _when(
                formula=["MOD(ROW(),2)=0"],
                fill=PatternFill(bgColor=ZEBRA, fill_type="solid"),
                border=ROW_BORDER,
            ),
        )
        sheet.conditional_formatting.add(data, _when(formula=["MOD(ROW(),2)=1"], border=ROW_BORDER))

        # A validação na cor do estado — o mesmo tom do ecrã.
        cells = f"{validation_column}{first_data}:{validation_column}{last_row}"
        top = f"{validation_column}{first_data}"
        for needle, state in (
            ("Confere", "match"),
            ("incorrectamente", "mismatch"),
            ("Não Creditado", "missing"),
            ("duplicados", "duplicated"),
            ("zerado", "zero"),
        ):
            _, color = STATE_TONES[state]
            sheet.conditional_formatting.add(
                cells,
                _when(
                    formula=[f'ISNUMBER(SEARCH("{needle}",{top}))'],
                    font=Font(bold=True, color=color),
                    stopIfTrue=True,
                ),
            )
        if difference_column:
            diff = f"{difference_column}{first_data}"
            sheet.conditional_formatting.add(
                f"{difference_column}{first_data}:{difference_column}{last_row}",
                _when(
                    formula=[f"AND(ISNUMBER({diff}),{diff}<>0)"],
                    font=Font(bold=True, color=STATE_TONES["mismatch"][1]),
                ),
            )

    # O POS ID fica à vista ao correr para o lado, e o cabeçalho ao correr para baixo.
    sheet.freeze_panes = f"C{first_data}"
    printable(sheet, repeat_rows=f"{header}:{header}")


def printable(sheet: Worksheet, *, repeat_rows: str | None = None) -> None:
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    sheet.print_options.horizontalCentered = True
    sheet.page_margins.left = sheet.page_margins.right = 0.4
    if repeat_rows:
        sheet.print_title_rows = repeat_rows
    footer = sheet.oddFooter
    if footer is not None and footer.left is not None and footer.right is not None:
        footer.left.text = "MozaOps · &A"
        footer.right.text = "Página &P de &N"


def _when(
    *,
    formula: list[str],
    fill: PatternFill | None = None,
    font: Font | None = None,
    border: Border | None = None,
    stopIfTrue: bool | None = None,
) -> Rule:
    """Uma regra de formatação condicional por fórmula."""
    return Rule(
        type="expression",
        formula=formula,
        dxf=DifferentialStyle(font=font, fill=fill, border=border),
        stopIfTrue=stopIfTrue,
    )


def ignore_number_as_text(data: bytes, ranges: dict[str, str]) -> bytes:
    """Desliga, nos intervalos dados, o aviso de «número guardado como texto».

    O openpyxl não escreve `<ignoredErrors>`, por isso acrescenta-se ao XML de cada
    folha depois de gravado. O elemento tem lugar fixo no esquema: antes de
    `drawing`, `legacyDrawing`, `tableParts` e `extLst`, ou no fim.
    """
    source = zipfile.ZipFile(BytesIO(data))
    workbook_xml = source.read("xl/workbook.xml").decode("utf-8")
    rels_xml = source.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    # Os atributos vêm por qualquer ordem: lê-se cada elemento e depois cada atributo.
    targets = {
        _attribute(tag, "Id"): _attribute(tag, "Target")
        for tag in re.findall(r"<Relationship\b[^>]*>", rels_xml)
    }
    files: dict[str, str] = {}
    for tag in re.findall(r"<sheet\b[^>]*>", workbook_xml):
        name, rel_id = _attribute(tag, "name"), _attribute(tag, "r:id")
        if name in ranges and rel_id in targets:
            files["xl/" + targets[rel_id].removeprefix("/xl/")] = ranges[name]

    output = BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename in files:
                xml = content.decode("utf-8")
                element = (
                    f'<ignoredErrors><ignoredError sqref="{files[item.filename]}" '
                    'numberStoredAsText="1"/></ignoredErrors>'
                )
                anchor = re.search(r"<(drawing|legacyDrawing|tableParts|extLst)\b", xml)
                at = anchor.start() if anchor else xml.rindex("</worksheet>")
                content = (xml[:at] + element + xml[at:]).encode("utf-8")
            target.writestr(item, content)
    return output.getvalue()


def _attribute(tag: str, name: str) -> str:
    match = re.search(rf'\s{re.escape(name)}="([^"]*)"', tag)
    return match.group(1) if match else ""
